"""Read both supplied Excel layouts. Never infer seconds from row counts."""
from __future__ import annotations

from dataclasses import dataclass, field, asdict
from datetime import datetime, time, timedelta
from pathlib import Path
import math
import re
import unicodedata

from openpyxl import load_workbook


@dataclass
class Row:
    sheet: str
    row: int
    transcript: str
    translation: str = ""
    gloss: str = ""
    group: str = "whole_audio"
    scope_start: float = 0.0
    scope_end: float | None = None
    start: float | None = None
    end: float | None = None
    alignment_text: str | None = None
    skip: bool = False
    manual: bool = False
    flags: list[str] = field(default_factory=list)

    @property
    def key(self):
        return f"{self.sheet}:{self.row}"

    def to_dict(self):
        return asdict(self)


def seconds(value, number_format="", *, excel_time=False):
    if value is None or value == "":
        return None
    if isinstance(value, timedelta):
        return value.total_seconds()
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError(f"Non-finite timestamp: {value}")
        # Excel serial fractions use a day only in time-formatted columns.
        is_time_format = bool(re.search(r"[hs]", number_format.lower()))
        return float(value) * (86400 if excel_time and is_time_format else 1)
    text = str(value).strip()
    if text.lower() in {"not applicable", "na", "n/a", "-", "none"}:
        return None
    if re.fullmatch(r"\d+(?::\d+){1,2}(?:\.\d+)?", text):
        total = 0.0
        for part in text.split(":"):
            total = 60 * total + float(part)
        return total
    try:
        result = float(text)
    except ValueError:
        raise ValueError(f"Invalid timestamp: {value!r}") from None
    if not math.isfinite(result):
        raise ValueError(f"Non-finite timestamp: {value}")
    return result


def header_key(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def read_rows(path: Path, duration: float, overrides=(), ignore_existing=False):
    book = load_workbook(path, read_only=True, data_only=True)
    result = []
    try:
        for sheet in book:
            cells = list(sheet.iter_rows())
            header_index = next((i for i, row in enumerate(cells[:20])
                                 if any(header_key(c.value) == "romanizedscript" for c in row)), None)
            if header_index is None:
                continue
            headers = {header_key(c.value): j for j, c in enumerate(cells[header_index]) if c.value}
            def cell(row, name):
                j = headers.get(name)
                return row[j] if j is not None and j < len(row) else None
            def value(row, name):
                c = cell(row, name)
                return c.value if c else None
            def stamp(row, name, excel_time=False):
                c = cell(row, name)
                return seconds(c.value, c.number_format, excel_time=excel_time) if c else None
            group, scope = "whole_audio", (0.0, duration)
            for i, cells_row in enumerate(cells[header_index + 1:], header_index + 2):
                text = value(cells_row, "romanizedscript")
                if text is None or not str(text).strip():
                    continue  # Blank lines and the pen-and-paper note are not utterances.
                try:
                    gs = stamp(cells_row, "starttimehhmmss", True)
                    ge = stamp(cells_row, "endtimehhmmss", True)
                    if (gs is None) != (ge is None):
                        raise ValueError("Coarse start/end must both be present")
                    if gs is not None:
                        if not 0 <= gs < ge <= duration + 1 / 16000:
                            raise ValueError(f"Invalid coarse window {gs}, {ge} for {duration}s audio")
                        group = str(value(cells_row, "audiofile") or f"group_{i}")
                        scope = (gs, min(ge, duration))
                    st = stamp(cells_row, "starttimeforsentences")
                    en = stamp(cells_row, "endtimeforsentences")
                    if "starttime" in headers:
                        st, en = stamp(cells_row, "starttime"), stamp(cells_row, "endtime")
                    if ignore_existing:
                        st = en = None
                    if (st is None) != (en is None):
                        raise ValueError("Sentence start/end must both be present")
                    result.append(Row(sheet.title, i, str(text),
                                      str(value(cells_row, "englishtranslation") or ""),
                                      str(value(cells_row, "wordlevelromanscriptenglishtranslation") or ""),
                                      group, *scope, st, en))
                except ValueError as e:
                    raise ValueError(f"{path.name}, {sheet.title}!{i}: {e}") from e
    finally:
        book.close()
    if not result:
        raise ValueError(f"No 'Romanized script' column with text found in {path}")
    index = {r.key: r for r in result}
    seen = set()
    for patch in overrides:
        key = f"{patch.get('sheet', 'Sheet1')}:{int(patch['row'])}"
        if key not in index:
            raise ValueError(f"Override references a missing transcript row: {key}")
        if key in seen:
            raise ValueError(f"Duplicate override: {key}")
        seen.add(key)
        r = index[key]
        if "alignment_text" in patch:
            r.alignment_text = str(patch["alignment_text"])
        r.skip = bool(patch.get("skip", False))
        for attr in ("scope_start", "scope_end", "start", "end"):
            if attr in patch and patch[attr] not in (None, ""):
                setattr(r, attr, seconds(patch[attr]))
        r.manual = "start" in patch or "end" in patch
        if r.manual and (r.start is None or r.end is None):
            raise ValueError(f"Manual override needs start and end: {key}")
        if not 0 <= r.scope_start < r.scope_end <= duration + 1 / 16000:
            raise ValueError(f"Invalid scope override: {key}")
        # A custom scope is a separate alignment block, even with the same group label.
    for r in result:
        if r.start is not None and not 0 <= r.scope_start <= r.start < r.end <= r.scope_end + 1 / 16000:
            raise ValueError(f"{r.key}: sentence times must be absolute seconds inside its coarse window")
    return result


def normalize(text: str):
    """Model-only text. Source spelling, translation and gloss are never rewritten."""
    flags = []
    if "/" in text:
        flags.append("alternative_or_slash_in_transcript")
    if "(" in text or ")" in text:
        flags.append("parenthetical_text_may_not_be_spoken")
    if re.search(r"\d", text):
        flags.append("numbers_need_spoken_form")
        text = re.sub(r"\b\w*\d\w*\b", " ", text)
    if text.strip().lower() in {"self introduction", "introduction"}:
        flags.append("possible_unspoken_heading")
    text = unicodedata.normalize("NFKD", text.replace("’", "'").replace("‘", "'"))
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    if any(c.isalpha() and not "a" <= c <= "z" for c in text):
        raise ValueError("Non-Latin letters found. Supply a romanized alignment_text override.")
    # Colon marks vowel length in these sheets: o:lh -> olh, not o lh.
    text = text.replace(":", "").replace("ː", "")
    text = re.sub(r"[^a-z'\s]", " ", text)
    words = [w.strip("'") for w in text.split() if w.strip("'")]
    if not words:
        raise ValueError("No alignable letters. Supply alignment_text or explicitly skip the row.")
    return words, flags


def find_inputs(folder: Path, config):
    def choose(setting, suffix):
        if config.get(setting):
            p = (folder / config[setting]).resolve()
            if not p.is_relative_to(folder.resolve()) or not p.is_file():
                raise ValueError(f"Invalid {setting} path under {folder}")
            return p
        files = sorted(p for p in folder.iterdir()
                       if p.is_file() and p.suffix.lower() == suffix and not p.name.startswith("~$"))
        if len(files) != 1:
            raise ValueError(f"Expected one root {suffix} in {folder}, found {len(files)}. Set '{setting}' in config.")
        return files[0]
    return choose("audio", ".wav"), choose("workbook", ".xlsx")
